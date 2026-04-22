import sys
import os
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from rosco.toolbox.ofTools.fast_io.output_processing import output_processing
from scipy.signal import butter, sosfiltfilt
from scipy.interpolate import interp1d
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '0_shared'))
import config as wt_config

def plot_not_available(ax, channel):
    return
    ax.set_ylabel(channel)
    ax.text(0.5, 0.5, 'not available', transform=ax.transAxes, ha='center')
    ax.grid(True)

def main():
    os.chdir(os.path.dirname(os.path.abspath(__file__)))
    

    channels = wt_config.CHANNELS_COMPARE


    ## 6-DOF
    u_rot = pd.read_csv(wt_config.FILTERED_WIND_CSV)
    resp_6 = pd.read_csv(wt_config.INTERP_6DOF_RESP_CSV)

    fig, axs = plt.subplots(len(channels), 1, sharex=True, figsize=(10, 2 * len(channels)), constrained_layout=True)
    

    axs[0].plot(resp_6['Time'], u_rot['U_avg(m/s)'], label='6-DOF Lookup', color='C0')
    axs[0].set_title('Average Wind Speed')
    for i, r in enumerate(channels[1:]):
        if r in resp_6.columns:
            axs[i+1].plot(resp_6['Time'], resp_6[r],color='C0',label='6-DOF Lookup')
            axs[i+1].set_ylabel(r)
        else:
            plot_not_available(axs[i+1], r)



    ## 1-DOF
    resp_1 = pd.read_csv(wt_config.SIM_1DOF_OUT_CSV)

    resp_1['GenPwr'] = resp_1['GenTq'] * resp_1['GenSpeed'] * 2 * np.pi / 60  # convert to kW

    for ax, channel in zip(axs, channels):
        if channel in resp_1.columns:
            ax.plot(resp_1['Time'], resp_1[channel], color='C1',label='6+1-DOF Sim',alpha=0.7)
            ax.set_ylabel(channel)
        else:
            plot_not_available(ax, channel)

    axs[-1].set_xlabel('Time (s)')

    
    
    ## OpenFAST output file
    filename = wt_config.OPENFAST_OUT_FILE

    op = output_processing()
    fastout = op.load_fast_out(filename)

    # fastout is a list; get the single dict
    fast_data = fastout[0]

    if True:
        dt = np.mean(np.diff(fast_data['Time']))
        sos = butter(wt_config.FILTER_ORDER, wt_config.FILTER_CUTOFF / 2.0, btype='low', fs=1.0 / dt, output='sos')
        rt_avg = sosfiltfilt(sos, fast_data['RtVAvgxh'])
        
        pd.DataFrame({'# Time(s)': fast_data['Time'], 'U_avg(m/s)': rt_avg}).to_csv(wt_config.FILTERED_WIND_CSV, index=False)


    for ax, channel in zip(axs, channels):
        if channel in fast_data:
            ax.plot(fast_data['Time'], fast_data[channel], color='C2',label='OpenFAST',alpha=0.7)
            unit_idx = fast_data['meta']['channels'].index(channel)
            unit = fast_data['meta']['attribute_units'][unit_idx]
            ax.set_ylabel(f'{channel}\n({unit})')
        else:
            plot_not_available(ax, channel)
    [a.grid() for a in axs]
    axs[-1].set_xlabel('Time (s)')
    axs[0].legend(loc='upper right')

    plt.tight_layout()
    axs[-1].set_xlabel('Time (s)')
    fig.align_ylabels()

    plt.show()

    ## Second figure: CHANNELS_6DOF only
    channels_6dof = wt_config.CHANNELS_6DOF
    fig2, axs2 = plt.subplots(len(channels_6dof), 1, sharex=True, figsize=(10, 2 * len(channels_6dof)), constrained_layout=True)

    for ax, ch in zip(axs2, channels_6dof):
        if ch in resp_6.columns:
            ax.plot(resp_6['Time'], resp_6[ch], color='C0', label='6-DOF Lookup')
        if ch in resp_1.columns:
            ax.plot(resp_1['Time'], resp_1[ch], color='C1', label='1-DOF Sim', alpha=0.7)
        if ch in fast_data:
            ax.plot(fast_data['Time'], fast_data[ch], color='C2', label='OpenFAST', alpha=0.7)
            unit_idx = fast_data['meta']['channels'].index(ch)
            unit = fast_data['meta']['attribute_units'][unit_idx]
            ax.set_ylabel(f'{ch}\n({unit})')
        else:
            ax.set_ylabel(ch)
        ax.grid(True)

    axs2[0].legend(loc='upper right')
    axs2[-1].set_xlabel('Time (s)')
    fig2.align_ylabels()

    plt.show()

    ## RMS error vs OpenFAST for t > 200 s
    t_min = 200.0
    t6 = resp_6['Time'].to_numpy()
    t1 = resp_1['Time'].to_numpy()
    tof = fast_data['Time']
    t_max = min(t6.max(), t1.max(), tof.max())

    mask_of = (tof >= t_min) & (tof <= t_max)
    t_common = tof[mask_of]

    print(f"\nRMS Error vs OpenFAST, t > {t_min} s  ({len(t_common)} points)")
    print(f"{'Channel':<14s}  {'6-DOF NRMS%':>12s}  {'1-DOF NRMS%':>12s}  {'Mean |OF|':>12s}")
    print("-" * 56)

    for ch in wt_config.CHANNELS_6DOF:
        if ch not in fast_data:
            continue
        yof = fast_data[ch][mask_of]
        mean_abs = np.mean(np.abs(yof))

        # 6-DOF lookup
        if ch in resp_6.columns:
            f6 = interp1d(t6, resp_6[ch].to_numpy(), kind='linear', bounds_error=False, fill_value='extrapolate')
            rms_6 = np.sqrt(np.mean((f6(t_common) - yof) ** 2))
            nrms_6 = 100.0 * rms_6 / mean_abs if mean_abs > 0 else float('inf')
        else:
            nrms_6 = float('nan')

        # 1-DOF sim
        if ch in resp_1.columns:
            f1 = interp1d(t1, resp_1[ch].to_numpy(), kind='linear', bounds_error=False, fill_value='extrapolate')
            rms_1 = np.sqrt(np.mean((f1(t_common) - yof) ** 2))
            nrms_1 = 100.0 * rms_1 / mean_abs if mean_abs > 0 else float('inf')
        else:
            nrms_1 = float('nan')

        print(f"{ch:<14s}  {nrms_6:12.1f}  {nrms_1:12.1f}  {mean_abs:12.2f}")

    # ── Build PowerPoint table ──
    prs = Presentation()
    prs.slide_width  = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # Title
    from pptx.enum.shapes import MSO_SHAPE
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(1.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(0x00, 0x5E, 0xA2)
    bar.line.fill.background()
    tf = bar.text_frame
    tf.margin_left = Inches(0.5)
    tf.margin_top  = Inches(0.15)
    p = tf.paragraphs[0]
    p.text = f"NRMS Error vs OpenFAST  (t > {t_min} s)"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Collect rows
    rows = []
    for ch in wt_config.CHANNELS_6DOF:
        if ch not in fast_data:
            continue
        yof = fast_data[ch][mask_of]
        mean_abs_ch = np.mean(np.abs(yof))
        n6 = 100.0 * np.sqrt(np.mean((interp1d(t6, resp_6[ch].to_numpy(), bounds_error=False, fill_value='extrapolate')(t_common) - yof) ** 2)) / mean_abs_ch if ch in resp_6.columns and mean_abs_ch > 0 else float('nan')
        n1 = 100.0 * np.sqrt(np.mean((interp1d(t1, resp_1[ch].to_numpy(), bounds_error=False, fill_value='extrapolate')(t_common) - yof) ** 2)) / mean_abs_ch if ch in resp_1.columns and mean_abs_ch > 0 else float('nan')
        rows.append((ch, n6, n1, mean_abs_ch))

    n_rows = len(rows) + 1  # +1 for header
    n_cols = 4
    tbl_w = Inches(10)
    tbl_h = Inches(0.45 * n_rows)
    left = (prs.slide_width - tbl_w) // 2
    top  = Inches(1.5)
    table = slide.shapes.add_table(n_rows, n_cols, left, top, tbl_w, tbl_h).table

    # Column widths
    table.columns[0].width = Inches(3.0)
    table.columns[1].width = Inches(2.5)
    table.columns[2].width = Inches(2.5)
    table.columns[3].width = Inches(2.0)

    # Header
    headers = ['Channel', '6-DOF NRMS (%)', '1-DOF NRMS (%)', 'Mean |OF|']
    for ci, h in enumerate(headers):
        cell = table.cell(0, ci)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x00, 0x5E, 0xA2)
        for par in cell.text_frame.paragraphs:
            par.font.size = Pt(14)
            par.font.bold = True
            par.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            par.alignment = PP_ALIGN.CENTER

    # Data rows
    for ri, (ch, n6, n1, mabs) in enumerate(rows, start=1):
        vals = [ch, f'{n6:.1f}', f'{n1:.1f}', f'{mabs:.2e}']
        for ci, v in enumerate(vals):
            cell = table.cell(ri, ci)
            cell.text = v
            if ri % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xE8, 0xF0, 0xF8)
            for par in cell.text_frame.paragraphs:
                par.font.size = Pt(13)
                par.alignment = PP_ALIGN.CENTER if ci > 0 else PP_ALIGN.LEFT

    pptx_path = os.path.join(wt_config.WAVE_TANK_DIR, 'rms_error_table.pptx')
    prs.save(pptx_path)
    print(f"\nSaved → {pptx_path}")


if __name__=="__main__":
    main()
